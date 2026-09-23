"""Text extraction for the document formats supported by Lexicon."""

from io import BytesIO
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}


class DocumentExtractionError(ValueError):
    """Raised when a supported document cannot be read into useful text."""


def extract_text(filename: str, file_bytes: bytes) -> tuple[str, int]:
    """Return readable text and a useful unit count for one supported document."""
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(file_bytes))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            unit_count = len(reader.pages)
        elif extension == ".docx":
            document = Document(BytesIO(file_bytes))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            text = "\n".join(paragraphs + table_cells)
            unit_count = 1
        elif extension == ".pptx":
            presentation = Presentation(BytesIO(file_bytes))
            text = "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text
            )
            unit_count = len(presentation.slides)
        elif extension == ".txt":
            text = file_bytes.decode("utf-8")
            unit_count = 1
        else:
            allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise DocumentExtractionError(f"Please upload one of these file types: {allowed}.")
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError(
            f"The uploaded {extension or 'file'} could not be read."
        ) from error

    text = text.strip()
    if not text:
        raise DocumentExtractionError(
            "No readable text was found. Image-only or scanned materials are not supported yet."
        )

    return text, unit_count
