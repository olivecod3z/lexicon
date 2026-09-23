import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pptx import Presentation

from app.main import app
from app.flashcards import FlashcardSet
from app.mcqs import MCQSet
from app.study_packs import StudyPack, generate_study_pack

client = TestClient(app)
TEST_DATABASE_PATH = Path(__file__).parent / ".main-materials-test.db"


def test_upload_once_then_generate_features_from_saved_material(monkeypatch) -> None:
    TEST_DATABASE_PATH.unlink(missing_ok=True)
    monkeypatch.setattr("app.materials.DATABASE_PATH", TEST_DATABASE_PATH)
    monkeypatch.setattr("app.main.generate_study_pack", lambda text: {"title": "Notes", "overview": "Overview", "learning_objectives": ["Learn one", "Learn two"], "sections": [{"heading": "One", "explanation": "First section", "key_points": ["Point"]}, {"heading": "Two", "explanation": "Second section", "key_points": ["Point"]}]})

    upload = client.post(
        "/materials",
        files={"file": ("lecture.txt", b"Cell biology lecture notes", "text/plain")},
    )

    assert upload.status_code == 201
    material_id = upload.json()["id"]
    assert upload.json()["filename"] == "lecture.txt"

    generated = client.post(f"/materials/{material_id}/study-pack")
    assert generated.status_code == 200
    assert generated.json()["title"] == "Notes"

    deleted = client.delete(f"/materials/{material_id}")
    assert deleted.status_code == 204

    missing = client.get(f"/materials/{material_id}")
    assert missing.status_code == 404

    TEST_DATABASE_PATH.unlink(missing_ok=True)


def test_swagger_shows_only_the_upload_once_feature_routes() -> None:
    paths = app.openapi()["paths"]

    assert "/materials/study-pack" not in paths
    assert "/materials/flashcards" not in paths
    assert "/materials/mcqs" not in paths
    assert "/materials/{material_id}/study-pack" in paths
    assert "/materials/{material_id}/flashcards" in paths
    assert "/materials/{material_id}/mcqs" in paths


def make_blank_pdf() -> bytes:
    """Create a tiny PDF with an empty page for endpoint validation tests."""
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_health_check_returns_ok() -> None:
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_root_redirects_to_api_docs() -> None:
    response = client.get("/", follow_redirects=False)

    assert response.status_code == 307
    assert response.headers["location"] == "/docs"


def test_extract_rejects_unsupported_filename() -> None:
    response = client.post(
        "/materials/extract",
        files={"file": ("notes.csv", b"lecture notes", "text/csv")},
    )

    assert response.status_code == 400
    assert ".pdf" in response.json()["detail"]


def test_extract_reports_pdf_without_selectable_text() -> None:
    response = client.post(
        "/materials/extract",
        files={"file": ("slides.pdf", make_blank_pdf(), "application/pdf")},
    )

    assert response.status_code == 422
    assert "No readable text was found" in response.json()["detail"]


def test_extracts_text_from_word_document() -> None:
    document = Document()
    document.add_paragraph("Cell biology lecture notes")
    output = BytesIO()
    document.save(output)

    response = client.post(
        "/materials/extract",
        files={
            "file": (
                "lecture.docx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["unit_count"] == 1
    assert response.json()["preview"] == "Cell biology lecture notes"


def test_extracts_text_from_plain_text_file() -> None:
    response = client.post(
        "/materials/extract",
        files={"file": ("lecture-notes.txt", b"Cell biology lecture notes", "text/plain")},
    )

    assert response.status_code == 200
    assert response.json()["unit_count"] == 1
    assert response.json()["preview"] == "Cell biology lecture notes"


def test_extracts_text_from_powerpoint() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(0, 0, 100, 100)
    text_box.text_frame.text = "Introduction to cells"
    output = BytesIO()
    presentation.save(output)

    response = client.post(
        "/materials/extract",
        files={
            "file": (
                "lecture.pptx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["unit_count"] == 1
    assert response.json()["preview"] == "Introduction to cells"


def test_study_pack_returns_validated_ai_result(monkeypatch) -> None:
    expected_pack = StudyPack(
        title="Cell Biology",
        overview="An introduction to cells and their major parts.",
        learning_objectives=["Identify cell structures.", "Explain basic cell functions."],
        sections=[
            {
                "heading": "Cells",
                "explanation": "Cells are the basic units described in the lecture.",
                "key_points": ["Cells have structures."],
            },
            {
                "heading": "Organelles",
                "explanation": "Organelles have distinct functions in the lecture.",
                "key_points": ["Functions differ by organelle."],
            },
        ],
    )
    monkeypatch.setattr("app.main.extract_text", lambda *_: ("lecture text", 1))
    monkeypatch.setattr("app.main.generate_study_pack", lambda _: expected_pack)

    response = client.post(
        "/materials/study-pack",
        files={"file": ("lecture.docx", b"not inspected in this test", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 200
    assert response.json() == expected_pack.model_dump()


def test_study_pack_reports_missing_ai_configuration(monkeypatch) -> None:
    monkeypatch.setattr("app.main.extract_text", lambda *_: ("lecture text", 1))
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("OPENAI_MODEL", raising=False)

    response = client.post(
        "/materials/study-pack",
        files={"file": ("lecture.pdf", b"not inspected in this test", "application/pdf")},
    )

    assert response.status_code == 503
    assert "not configured" in response.json()["detail"]


def test_flashcard_endpoint_returns_validated_cards(monkeypatch) -> None:
    expected_set = FlashcardSet(
        title="Cell Biology Flashcards",
        flashcards=[
            {"question": "What is a cell?", "answer": "A basic unit of life.", "topic": "Cells", "card_type": "question_answer"},
            {"question": "An organelle has a ____.", "answer": "function", "topic": "Organelles", "card_type": "fill_in_the_blank"},
            {"question": "What is a nucleus?", "answer": "A cell structure.", "topic": "Organelles", "card_type": "question_answer"},
            {"question": "Cells contain ____.", "answer": "structures", "topic": "Cells", "card_type": "fill_in_the_blank"},
            {"question": "What do organelles do?", "answer": "Functions.", "topic": "Organelles", "card_type": "question_answer"},
        ],
    )
    monkeypatch.setattr("app.main.extract_text", lambda *_: ("lecture text", 1))
    monkeypatch.setattr("app.main.generate_flashcards", lambda _: expected_set)

    response = client.post(
        "/materials/flashcards",
        files={"file": ("lecture.txt", b"not inspected in this test", "text/plain")},
    )

    assert response.status_code == 200
    assert response.json() == expected_set.model_dump()


def test_mcq_endpoint_returns_validated_questions(monkeypatch) -> None:
    questions = [
        {
            "question": f"What is concept {number}?",
            "options": [
                {"label": "A", "text": "The correct definition."},
                {"label": "B", "text": "A different idea."},
                {"label": "C", "text": "An unrelated process."},
                {"label": "D", "text": "A false statement."},
            ],
            "correct_option": "A",
            "explanation": "The lecture gives this definition.",
            "topic": "Concepts",
        }
        for number in range(1, 6)
    ]
    expected_set = MCQSet(title="Lecture MCQs", questions=questions)
    monkeypatch.setattr("app.main.extract_text", lambda *_: ("lecture text", 1))
    monkeypatch.setattr("app.main.generate_mcqs", lambda _: expected_set)

    response = client.post(
        "/materials/mcqs",
        files={"file": ("lecture.txt", b"not inspected in this test", "text/plain")},
    )

    assert response.status_code == 200
    assert response.json() == expected_set.model_dump()


def test_generation_requests_strict_json_without_storing_response(monkeypatch) -> None:
    expected_pack = StudyPack(
        title="Cell Biology",
        overview="An introduction to cells and their major parts.",
        learning_objectives=["Identify cell structures.", "Explain basic cell functions."],
        sections=[
            {
                "heading": "Cells",
                "explanation": "Cells are the basic units described in the lecture.",
                "key_points": ["Cells have structures."],
            },
            {
                "heading": "Organelles",
                "explanation": "Organelles have distinct functions in the lecture.",
                "key_points": ["Functions differ by organelle."],
            },
        ],
    )
    captured_request = {}

    class FakeResponses:
        def create(self, **kwargs):
            captured_request.update(kwargs)
            return SimpleNamespace(output_text=json.dumps(expected_pack.model_dump()))

    class FakeClient:
        responses = FakeResponses()

    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setenv("OPENAI_MODEL", "test-model")
    monkeypatch.setattr("app.study_packs.OpenAI", lambda api_key: FakeClient())

    result = generate_study_pack("Lecture source text")

    assert result == expected_pack
    assert captured_request["store"] is False
    assert captured_request["text"]["format"]["type"] == "json_schema"
    assert captured_request["text"]["format"]["strict"] is True
